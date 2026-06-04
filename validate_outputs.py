from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from code_partie1_stucture_generale.page_fields import find_numeric_zones
from code_partie1_stucture_generale.pdf_utils import pdf_to_images, correct_orientation
from code_partie2_elements_graphiques import read_exam_choices, read_page01_graphics
from code_partie2_elements_graphiques.page01 import GROUP_ROI, SIGNATURE_FALLBACK_ROI, STUDENT_ID_ROI
from config import DEBUG_DIR, PROJECT_ROOT
from data_discovery import IMAGE_SUFFIXES, discover_datasets


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate real Program 1/2 outputs and generate review debug images.")
    parser.add_argument("results_folder", nargs="?", help="Folder containing generated .xlsx files, e.g. FORM1_RESULTS")
    parser.add_argument("--exam-root", default=None, help="Real data folder containing presence images and PDFs")
    parser.add_argument("--debug-root", default=str(DEBUG_DIR), help="Debug root folder")
    parser.add_argument("--presence-samples", type=int, default=10)
    parser.add_argument("--pdf-samples", type=int, default=5)
    parser.add_argument("--pages-per-pdf", type=int, default=2)
    parser.add_argument("--discover", action="store_true", help="List detected datasets and exit")
    args = parser.parse_args()

    if args.discover:
        for dataset in discover_datasets(PROJECT_ROOT):
            print(
                f"{dataset.exam_name}: presences={dataset.presences_folder} images={dataset.image_count}, "
                f"pdfs={dataset.pdf_count}, signatures={dataset.signature_count}, results={dataset.results_folder}"
            )
        return

    if not args.results_folder:
        raise SystemExit("results_folder is required unless --discover is used")

    results_folder = Path(args.results_folder).resolve()
    exam_root = _infer_exam_root(results_folder, args.exam_root)
    debug_root = Path(args.debug_root).resolve()

    summary = validate_results(results_folder, exam_root)
    review_summary = generate_review_images(
        exam_root=exam_root,
        debug_root=debug_root,
        presence_samples=args.presence_samples,
        pdf_samples=args.pdf_samples,
        pages_per_pdf=args.pages_per_pdf,
    )
    summary["review"] = review_summary
    summary["qcm_debug"] = _qcm_debug_summary(debug_root / "qcm" / "qcm_candidates.csv")
    summary["handwritten_debug"] = _csv_status_summary(debug_root / "handwritten_digits" / "handwritten_debug.csv")
    summary["signature_debug"] = _csv_status_summary(debug_root / "signatures" / "signature_scores.csv")
    summary["printed_ocr_debug"] = _csv_status_summary(debug_root / "page01" / "printed_ocr_debug.csv")

    results_folder.mkdir(parents=True, exist_ok=True)
    summary_path = results_folder / "validation_summary.json"
    summary_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")

    print(json.dumps(summary, indent=2, ensure_ascii=False))
    print(f"Validation summary saved to {summary_path}")


def validate_results(results_folder: Path, exam_root: Path | None) -> dict[str, Any]:
    images = _list_images(exam_root) if exam_root else []
    pdfs = _list_pdfs(exam_root) if exam_root else []
    summary: dict[str, Any] = {
        "results_folder": str(results_folder),
        "exam_root": str(exam_root) if exam_root else "",
        "program1": {
            "total_presence_images": len(images),
            "presence_excel": "",
            "rows_in_presence_excel": 0,
            "columns_ok": False,
            "row_count_ok": False,
            "missing_image_names": [],
            "status": "missing",
        },
        "program2": {
            "total_pdfs": len(pdfs),
            "generated_excel_count": 0,
            "missing_excel_count": 0,
            "invalid_excel_count": 0,
            "missing_excels": [],
            "invalid_excels": [],
        },
    }

    presence_files = sorted(results_folder.glob("*PRESENCES.xlsx"))
    if presence_files:
        summary["program1"].update(_validate_presence_excel(presence_files[0], images))

    generated = 0
    missing: list[str] = []
    invalid: list[dict[str, str]] = []
    for pdf in pdfs:
        xlsx = results_folder / f"{pdf.stem}.xlsx"
        if not xlsx.exists():
            missing.append(pdf.name)
            continue
        generated += 1
        reason = _validate_form_workbook(xlsx)
        if reason:
            invalid.append({"file": xlsx.name, "reason": reason})
    summary["program2"].update(
        {
            "generated_excel_count": generated,
            "missing_excel_count": len(missing),
            "invalid_excel_count": len(invalid),
            "missing_excels": missing,
            "invalid_excels": invalid,
        }
    )
    return summary


def generate_review_images(
    exam_root: Path | None,
    debug_root: Path,
    presence_samples: int,
    pdf_samples: int,
    pages_per_pdf: int,
) -> dict[str, Any]:
    review_root = debug_root / "review"
    page01_dir = review_root / "page01_alignment"
    signatures_dir = review_root / "signatures"
    qcm_dir = review_root / "qcm"
    handwritten_dir = review_root / "handwritten"
    for folder in [page01_dir, signatures_dir, qcm_dir, handwritten_dir]:
        folder.mkdir(parents=True, exist_ok=True)

    if exam_root is None:
        return {"status": "skipped", "reason": "no exam root"}

    saved = {
        "presence_images_reviewed": 0,
        "pdfs_reviewed": 0,
        "pdf_pages_reviewed": 0,
        "page01_alignment_images": 0,
        "signature_crops": 0,
        "qcm_overlays": 0,
        "handwritten_crops": 0,
        "errors": [],
    }

    image_attempts = 0
    for image_path in _list_images(exam_root):
        if saved["presence_images_reviewed"] >= presence_samples:
            break
        image_attempts += 1
        try:
            image = correct_orientation(Image.open(image_path).convert("RGB"))
            graphics = read_page01_graphics(image)
            _save_page01_overlay(image, graphics, page01_dir / f"{image_path.stem}_page01.png")
            graphics.signature_crop.save(signatures_dir / f"{image_path.stem}_signature.png")
            _save_relative_crop(image, STUDENT_ID_ROI, page01_dir / f"{image_path.stem}_student_grid.png")
            saved["presence_images_reviewed"] += 1
            saved["page01_alignment_images"] += 1
            saved["signature_crops"] += 1
        except Exception as exc:
            saved["errors"].append({"file": image_path.name, "reason": str(exc)})
    saved["presence_images_attempted"] = image_attempts

    for pdf_path in _list_pdfs(exam_root)[:pdf_samples]:
        try:
            pages = [correct_orientation(img) for img in pdf_to_images(str(pdf_path))]
        except Exception as exc:
            saved["errors"].append({"file": pdf_path.name, "reason": str(exc)})
            continue
        if not pages:
            saved["errors"].append({"file": pdf_path.name, "reason": "no rendered pages"})
            continue
        saved["pdfs_reviewed"] += 1
        page_indices = [0]
        page_indices.extend(range(4, min(len(pages), 4 + max(1, pages_per_pdf))))
        if len(page_indices) < 2 and len(pages) > 1:
            page_indices.append(1)

        for page_index in page_indices:
            page = pages[page_index]
            page_no = page_index + 1
            saved["pdf_pages_reviewed"] += 1
            if page_index == 0:
                try:
                    graphics = read_page01_graphics(page)
                    _save_page01_overlay(page, graphics, page01_dir / f"{pdf_path.stem}_p{page_no}_page01.png")
                    graphics.signature_crop.save(signatures_dir / f"{pdf_path.stem}_p{page_no}_signature.png")
                    saved["page01_alignment_images"] += 1
                    saved["signature_crops"] += 1
                except Exception as exc:
                    saved["errors"].append({"file": f"{pdf_path.name}:p{page_no}", "reason": str(exc)})
                continue

            try:
                questions = read_exam_choices(
                    page,
                    debug_dir=str(debug_root / "qcm"),
                    source_name=f"review_{pdf_path.stem}_p{page_no}",
                )
                _save_qcm_overlay(page, questions, qcm_dir / f"{pdf_path.stem}_p{page_no}_qcm.png")
                saved["qcm_overlays"] += 1
            except Exception as exc:
                saved["errors"].append({"file": f"{pdf_path.name}:p{page_no}", "reason": f"qcm:{exc}"})

            try:
                zones = find_numeric_zones(page)
                for zone_index, zone in enumerate(zones):
                    for label in ["mantisse", "exposant"]:
                        crop = zone.get(label)
                        if crop is None:
                            continue
                        path = handwritten_dir / f"{pdf_path.stem}_p{page_no}_z{zone_index}_{label}_{_crop_quality(crop)}.png"
                        crop.save(path)
                        saved["handwritten_crops"] += 1
            except Exception as exc:
                saved["errors"].append({"file": f"{pdf_path.name}:p{page_no}", "reason": f"handwritten:{exc}"})
    return saved


def _validate_presence_excel(path: Path, images: list[Path]) -> dict[str, Any]:
    expected_headers = ["imageName", "studentID_grid", "studentID_signature"]
    result = {
        "presence_excel": str(path),
        "rows_in_presence_excel": 0,
        "columns_ok": False,
        "row_count_ok": False,
        "missing_image_names": [],
        "status": "invalid",
    }
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        headers = [ws.cell(row=1, column=i).value for i in range(1, 4)]
        result["columns_ok"] = headers == expected_headers
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(cell is not None for cell in row):
                rows.append(row)
        result["rows_in_presence_excel"] = len(rows)
        result["row_count_ok"] = len(rows) == len(images)
        input_names = {p.name for p in images}
        result["missing_image_names"] = [row[0] for row in rows if row and row[0] and row[0] not in input_names]
        result["status"] = "ok" if result["columns_ok"] and result["row_count_ok"] and not result["missing_image_names"] else "invalid"
        wb.close()
    except Exception as exc:
        result["error"] = str(exc)
    return result


def _validate_form_workbook(path: Path) -> str:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        missing = [sheet for sheet in ["PAGE-01", "EXAM"] if sheet not in wb.sheetnames]
        wb.close()
        return f"missing sheets: {', '.join(missing)}" if missing else ""
    except Exception as exc:
        return str(exc)


def _save_page01_overlay(image: Image.Image, graphics: Any, path: Path) -> None:
    scale = 900 / max(1, image.width)
    preview = image.resize((900, max(1, int(image.height * scale))))
    draw = ImageDraw.Draw(preview)
    _draw_relative_box(draw, image.size, preview.size, STUDENT_ID_ROI, "student_id", "blue")
    _draw_relative_box(draw, image.size, preview.size, GROUP_ROI, "group", "purple")
    _draw_relative_box(draw, image.size, preview.size, SIGNATURE_FALLBACK_ROI, "signature_roi", "green")
    for box in getattr(graphics.student_grid, "boxes", {}).values():
        _draw_box(draw, image.size, preview.size, box, "cyan")
    for box in getattr(graphics.group_grid, "boxes", {}).values():
        _draw_box(draw, image.size, preview.size, box, "magenta")
    path.parent.mkdir(parents=True, exist_ok=True)
    preview.save(path)


def _save_qcm_overlay(image: Image.Image, questions: list[Any], path: Path) -> None:
    scale = 900 / max(1, image.width)
    preview = image.resize((900, max(1, int(image.height * scale))))
    draw = ImageDraw.Draw(preview)
    for q in questions:
        for label, box in q.boxes.items():
            checked = q.choices.get(label, 0) == 1
            _draw_box(draw, image.size, preview.size, box, "green" if checked else "orange")
    path.parent.mkdir(parents=True, exist_ok=True)
    preview.save(path)


def _draw_relative_box(
    draw: ImageDraw.ImageDraw,
    original_size: tuple[int, int],
    preview_size: tuple[int, int],
    roi: tuple[float, float, float, float],
    label: str,
    color: str,
) -> None:
    width, height = original_size
    x0, y0, x1, y1 = roi
    box = (int(x0 * width), int(y0 * height), int(x1 * width), int(y1 * height))
    _draw_abs_box(draw, original_size, preview_size, box, label, color)


def _draw_box(
    draw: ImageDraw.ImageDraw,
    original_size: tuple[int, int],
    preview_size: tuple[int, int],
    box: Any,
    color: str,
) -> None:
    _draw_abs_box(draw, original_size, preview_size, (box.x0, box.y0, box.x1, box.y1), "", color)


def _draw_abs_box(
    draw: ImageDraw.ImageDraw,
    original_size: tuple[int, int],
    preview_size: tuple[int, int],
    box: tuple[int, int, int, int],
    label: str,
    color: str,
) -> None:
    sx = preview_size[0] / max(1, original_size[0])
    sy = preview_size[1] / max(1, original_size[1])
    xy = tuple(int(v) for v in (box[0] * sx, box[1] * sy, box[2] * sx, box[3] * sy))
    draw.rectangle(xy, outline=color, width=2)
    if label:
        draw.text((xy[0] + 3, xy[1] + 3), label, fill=color)


def _save_relative_crop(image: Image.Image, roi: tuple[float, float, float, float], path: Path) -> None:
    width, height = image.size
    crop = image.crop((int(roi[0] * width), int(roi[1] * height), int(roi[2] * width), int(roi[3] * height)))
    path.parent.mkdir(parents=True, exist_ok=True)
    crop.save(path)


def _qcm_debug_summary(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"csv": str(path), "rows": 0, "accepted": 0, "rejected": 0, "reasons": {}}
    reasons: dict[str, int] = {}
    accepted = 0
    rows = 0
    with path.open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            rows += 1
            reason = row.get("reason", "")
            reasons[reason] = reasons.get(reason, 0) + 1
            accepted += 1 if str(row.get("accepted", "0")) == "1" else 0
    return {"csv": str(path), "rows": rows, "accepted": accepted, "rejected": rows - accepted, "reasons": reasons}


def _csv_status_summary(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"csv": str(path), "rows": 0, "statuses": {}}
    statuses: dict[str, int] = {}
    rows = 0
    with path.open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            rows += 1
            status = row.get("status", "")
            statuses[status] = statuses.get(status, 0) + 1
    return {"csv": str(path), "rows": rows, "statuses": statuses}


def _crop_quality(crop: Image.Image | np.ndarray | None) -> str:
    if crop is None:
        return "invalid"
    arr = np.asarray(crop)
    if arr.size == 0:
        return "empty_crop"
    if arr.ndim == 3:
        arr = arr[..., :3].mean(axis=2)
    if min(arr.shape[:2]) < 8:
        return "too_small"
    arr = np.nan_to_num(arr, nan=255.0, posinf=255.0, neginf=0.0)
    contrast = float(arr.max() - arr.min())
    dark_ratio = float((arr < 150).mean())
    border = max(1, min(arr.shape[:2]) // 20)
    border_pixels = np.concatenate([arr[:border, :].ravel(), arr[-border:, :].ravel(), arr[:, :border].ravel(), arr[:, -border:].ravel()])
    border_dark = float((border_pixels < 150).mean())
    if dark_ratio < 0.003:
        return "empty_crop"
    if contrast < 25:
        return "low_contrast"
    if border_dark > max(0.10, dark_ratio * 2.5):
        return "border_contamination"
    if dark_ratio > 0.55:
        return "too_noisy"
    return "valid_crop"


def _infer_exam_root(results_folder: Path, explicit_exam_root: str | None) -> Path | None:
    if explicit_exam_root:
        return Path(explicit_exam_root).resolve()
    name = results_folder.name
    if name.endswith("_RESULTS"):
        candidate = results_folder.parent / name[: -len("_RESULTS")]
        if candidate.exists():
            return candidate
    return None


def _list_images(folder: Path | None) -> list[Path]:
    if folder is None or not folder.exists():
        return []
    return sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in IMAGE_SUFFIXES)


def _list_pdfs(folder: Path | None) -> list[Path]:
    if folder is None or not folder.exists():
        return []
    return sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() == ".pdf")


if __name__ == "__main__":
    main()
