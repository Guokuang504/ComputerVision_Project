from __future__ import annotations

import argparse
import csv
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


def main() -> None:
    parser = argparse.ArgumentParser(description="Build an optional Student ID -> name roster.")
    parser.add_argument(
        "--forms",
        nargs="+",
        default=["FORM1", "FORM2", "FORM3"],
        help="Exam form folders containing reference Excel files.",
    )
    parser.add_argument(
        "--output",
        default="annotations/student_roster.csv",
        help="CSV output path used by the main pipeline.",
    )
    args = parser.parse_args()

    votes: dict[str, Counter[tuple[str, str]]] = defaultdict(Counter)
    skipped: list[tuple[str, str]] = []

    for form_dir in args.forms:
        for workbook_path in sorted(Path(form_dir).glob("EXAM_*.xlsx")):
            try:
                wb = openpyxl.load_workbook(workbook_path, data_only=True)
                ws = wb["PAGE-01"]
            except Exception as exc:
                skipped.append((str(workbook_path), type(exc).__name__))
                continue

            student_id = _clean_id(ws["B17"].value)
            firstname = _clean_name(ws["B13"].value)
            lastname = _clean_name(ws["B14"].value)
            if student_id and (firstname or lastname):
                votes[student_id][(firstname, lastname)] += 1

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["student_id", "firstname", "lastname"])
        writer.writeheader()
        for student_id in sorted(votes, key=lambda value: int(value) if value.isdigit() else value):
            (firstname, lastname), _count = votes[student_id].most_common(1)[0]
            writer.writerow({
                "student_id": student_id,
                "firstname": firstname,
                "lastname": lastname,
            })

    print(f"Wrote {len(votes)} roster entries to {output_path}")
    if skipped:
        print(f"Skipped {len(skipped)} invalid workbook(s)")


def _clean_id(value) -> str:
    if value is None:
        return ""
    try:
        return str(int(float(str(value).strip())))
    except (TypeError, ValueError):
        return "".join(ch for ch in str(value) if ch.isdigit())


def _clean_name(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


if __name__ == "__main__":
    main()
