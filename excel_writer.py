from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook


PAGE01_ROWS = [
    ("Module", "module"),
    ("Professor", "professor"),
    ("Date", "date"),
    ("Code", "code"),
    ("Notes de cours", "notes_de_cours"),
    ("Notes manuscrites", "notes_manuscrites"),
    ("Ordinateur portable", "ordinateur_portable"),
    ("Calculatrice ", "calculatrice"),
    ("Feuilles brouillon", "feuilles_brouillon"),
    ("Note maximale", "note_max"),
    ("Note pour valider", "note_valid"),
    ("", None),
    ("Prénom", "firstname"),
    ("Nom", "lastname"),
    ("Validation signature", "validation_signature"),
    ("Group", "group"),
    ("STUDENT ID", "student_id"),
    ("Validation cryptogramme", "validation_cryptogramme"),
]

EXAM_HEADERS = [
    "QUESTION",
    "CHOIX A",
    "CHOIX B",
    "CHOIX C",
    "CHOIX D",
    "CHOIX E",
    "CHOIX F",
    "CHOIX G",
    "CHOIX H",
    "MANTISSE",
    "EXPOSANT",
    "UNITE",
]


def write_presences_xlsx(rows: list[dict[str, Any]], output_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PRESENCES"
    headers = ["imageName", "studentID_grid", "studentID_signature"]
    ws.append(headers)
    for row in rows:
        ws.append([_clean(row.get(key, "")) for key in headers])
    _save(wb, output_path)


def write_form_xlsx(page01_data: dict[str, Any], exam_data: list[dict[str, Any]], output_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PAGE-01"
    for row_index, (label, key) in enumerate(PAGE01_ROWS, start=1):
        ws.cell(row=row_index, column=1, value=label)
        if key is not None:
            ws.cell(row=row_index, column=2, value=_clean(page01_data.get(key, "")))

    exam_ws = wb.create_sheet("EXAM")
    exam_ws.append(EXAM_HEADERS)
    for row in exam_data:
        exam_ws.append(
            [
                _clean(row.get("question", "")),
                _choice(row.get("choix_A")),
                _choice(row.get("choix_B")),
                _choice(row.get("choix_C")),
                _choice(row.get("choix_D")),
                _choice(row.get("choix_E")),
                _choice(row.get("choix_F")),
                _choice(row.get("choix_G")),
                _choice(row.get("choix_H")),
                _clean(row.get("mantisse", "")),
                _clean(row.get("exposant", "")),
                _clean(row.get("unite", "")),
            ]
        )
    _save(wb, output_path)


def _choice(value: Any) -> str:
    return "X" if value == 1 else ""


def _clean(value: Any) -> Any:
    if value is None:
        return ""
    return value


def _save(wb: Workbook, output_path: str) -> None:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
